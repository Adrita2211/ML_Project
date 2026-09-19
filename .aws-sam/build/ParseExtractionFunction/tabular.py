"""
Reads CSV or XLSX uploads into a list of dicts with camelCase keys, so the
same ingestion code accepts an MMIT-style export (snake_case columns such as
actual_tier) and hand-written files (actualTier) alike.

Values are normalized: blanks become None, dates become ISO strings, and
whole-number floats from Excel become ints. Type-specific parsing (tier
numbers, Y/N flags) stays with the consumers.
"""
import csv
import io
import re
from datetime import date, datetime


def camelize(key: str) -> str:
    parts = [p for p in re.split(r"[_\s]+", str(key).strip()) if p]
    if not parts:
        return ""
    return parts[0][:1].lower() + parts[0][1:] + "".join(p[:1].upper() + p[1:] for p in parts[1:])


def _clean(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def read_rows(body: bytes, filename: str, sheet: str | None = None) -> list[dict]:
    """Rows of a .csv or .xlsx file. For .xlsx, `sheet` picks a worksheet by
    name and falls back to the first sheet when absent."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl  # bundled via lambdas/requirements.txt

        workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        worksheet = workbook[sheet] if sheet and sheet in workbook.sheetnames else workbook.worksheets[0]
        raw_rows = list(worksheet.iter_rows(values_only=True))
    else:
        text = body.decode("utf-8-sig")
        raw_rows = [tuple(r) for r in csv.reader(io.StringIO(text))]

    if not raw_rows:
        return []

    headers = [camelize(h) if h is not None else "" for h in raw_rows[0]]
    rows = []
    for raw in raw_rows[1:]:
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in raw):
            continue
        rows.append({h: _clean(v) for h, v in zip(headers, raw) if h})
    return rows
