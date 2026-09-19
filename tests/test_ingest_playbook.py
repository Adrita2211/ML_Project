"""
Tests for playbook ingestion (lambdas/ingest_playbook.py): the Excel and CSV
versions of the playbook produce the same vector documents, and those documents
line up with what the rules and the retrieval filter expect. No AWS access needed.

The alignment tests matter most. Retrieval filters on the exact classification
string the rules emit. A playbook row spelled differently is not an error
anywhere; it simply never gets retrieved, and the case gets no actions.
"""
import io
import json
import os
import sys
import unittest
from unittest import mock

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "lambdas"))
os.environ.setdefault("AWS_DEFAULT_REGION", "us-east-1")

with mock.patch("boto3.client"):
    import ingest_playbook
import deviation_rules  # noqa: E402

try:
    import yaml
except ImportError:
    yaml = None

SAMPLE = os.path.join(ROOT, "sample_data")
XLSX = "Persona_NBA_Playbook.xlsx"
CSV = "persona_nba_playbook.csv"


def ingest(filename):
    """Run the real handler on a sample file; return (documents it would embed, result)."""
    with open(os.path.join(SAMPLE, filename), "rb") as f:
        body = f.read()
    s3 = mock.Mock()
    s3.get_object.return_value = {"Body": io.BytesIO(body)}
    ingest_playbook.s3 = s3
    captured = []
    with mock.patch.object(ingest_playbook.vector_store, "put_documents",
                           side_effect=lambda docs: captured.extend(docs) or len(docs)):
        result = ingest_playbook.handler(
            {"detail": {"bucket": {"name": "b"}, "object": {"key": f"playbook/{filename}"}}}, None)
    return captured, result


class Ingestion(unittest.TestCase):
    def test_the_excel_playbook_is_ingested_row_for_row(self):
        docs, result = ingest(XLSX)
        self.assertEqual((result["indexed"], result["skippedRows"]), (7, 0))
        self.assertEqual(len(docs), 7)
        self.assertEqual(len({d["key"] for d in docs}), 7, "every row needs its own key, or one overwrites another")

    def test_excel_and_csv_give_identical_vector_documents(self):
        def content(docs):  # sourceKey records which file a row came from, so it legitimately differs
            return [{**d, "metadata": {k: v for k, v in d["metadata"].items() if k != "sourceKey"}} for d in docs]

        from_xlsx, _ = ingest(XLSX)
        from_csv, _ = ingest(CSV)
        self.assertEqual(content(from_xlsx), content(from_csv), "the two copies of the playbook have drifted apart")

    def test_a_document_carries_what_retrieval_and_the_validator_need(self):
        docs, _ = ingest(XLSX)
        doc = next(d for d in docs if d["key"] == "playbook#tier-downgrade#frm")
        meta = doc["metadata"]
        self.assertEqual((meta["docType"], meta["classification"], meta["persona"]), ("playbook", "Tier downgrade", "FRM"))
        self.assertEqual((meta["approvalRequired"], meta["deliveryChannel"]), ("Y", "CRM task"))
        self.assertEqual(meta["doNotDo"], "Do not negotiate contract terms")
        self.assertIn("approved current-access guidance", meta["nextBestAction"])
        # what gets embedded is the meaning of the row, so similarity search has something to match on
        self.assertIn("Tier downgrade", doc["embedText"])
        self.assertIn("Identify priority HCP accounts", doc["embedText"])

    def test_the_data_quality_row_is_normalized_and_needs_no_approval(self):
        docs, _ = ingest(XLSX)
        row = next(d for d in docs if d["metadata"]["persona"] == "Data Steward")
        self.assertEqual(row["metadata"]["classification"], "Data-quality exception")  # source says "Any data-quality exception"
        self.assertEqual(row["metadata"]["approvalRequired"], "N")

    def test_re_ingesting_reuses_the_same_keys_so_edits_overwrite_not_duplicate(self):
        first, _ = ingest(XLSX)
        second, _ = ingest(XLSX)
        self.assertEqual([d["key"] for d in first], [d["key"] for d in second])

    def test_the_workbook_the_business_edits_is_read_by_sheet_name_or_first_sheet(self):
        import openpyxl
        wb = openpyxl.load_workbook(os.path.join(SAMPLE, XLSX))
        self.assertEqual(wb.sheetnames[0], "Persona_NBA_Playbook")

    def test_a_workbook_missing_required_columns_is_rejected_loudly(self):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(["persona", "something_else"])
        wb.active.append(["FRM", "x"])
        buffer = io.BytesIO()
        wb.save(buffer)
        s3 = mock.Mock()
        s3.get_object.return_value = {"Body": io.BytesIO(buffer.getvalue())}
        ingest_playbook.s3 = s3
        with self.assertRaises(ValueError) as caught:
            ingest_playbook.handler({"detail": {"bucket": {"name": "b"}, "object": {"key": "playbook/bad.xlsx"}}}, None)
        self.assertIn("classification", str(caught.exception))

    def test_the_original_workbook_layout_with_an_unnamed_first_column_still_works(self):
        """The source workbook's classification column was headed 'Column1'."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Persona_NBA_Playbook"
        ws.append(["Column1", "persona", "trigger_condition", "next_best_action", "required_evidence",
                   "approval_required", "delivery_channel", "do_not_do"])
        ws.append(["Coverage loss", "National Account Director", "t", "Add issue to agenda.", "e", "Y", "Account brief", "d"])
        buffer = io.BytesIO()
        wb.save(buffer)
        s3 = mock.Mock()
        s3.get_object.return_value = {"Body": io.BytesIO(buffer.getvalue())}
        ingest_playbook.s3 = s3
        with mock.patch.object(ingest_playbook.vector_store, "put_documents", side_effect=lambda d: len(d)) as put:
            ingest_playbook.handler({"detail": {"bucket": {"name": "b"}, "object": {"key": "playbook/orig.xlsx"}}}, None)
        self.assertEqual(put.call_args[0][0][0]["metadata"]["classification"], "Coverage loss")


class AlignmentWithTheRules(unittest.TestCase):
    def setUp(self):
        self.docs, _ = ingest(XLSX)
        self.classifications = {d["metadata"]["classification"] for d in self.docs}

    def test_every_playbook_classification_is_one_the_rules_can_emit(self):
        self.assertEqual(sorted(self.classifications - set(deviation_rules.CLASSIFICATION_ROUTING)), [],
                         "these would never be retrieved: no deviation is ever classified that way")

    def test_every_classification_the_rules_can_emit_has_at_least_one_action(self):
        self.assertEqual(sorted(set(deviation_rules.CLASSIFICATION_ROUTING) - self.classifications), [],
                         "a deviation of this kind would reach a person with no next action")

    def test_every_row_has_the_fields_that_drive_the_case(self):
        for doc in self.docs:
            meta = doc["metadata"]
            with self.subTest(row=doc["key"]):
                self.assertIn(meta["approvalRequired"], ("Y", "N"))
                for field in ("persona", "nextBestAction", "deliveryChannel", "doNotDo", "requiredEvidence"):
                    self.assertTrue(meta[field], f"{field} is empty")

    def test_the_retrieval_filter_finds_the_rows_the_case_workflow_asks_for(self):
        """Mimic case_logic.query_playbook against these documents, per classification."""
        import case_logic
        by_class = {}
        for d in self.docs:
            by_class.setdefault(d["metadata"]["classification"], []).append(d)

        def fake_query(text, metadata_filter, top_k):
            wanted = {list(c)[0]: list(c.values())[0]["$eq"] for c in metadata_filter["$and"]}
            return [{"key": d["key"], "metadata": d["metadata"]} for d in by_class.get(wanted["classification"], [])]

        for classification in deviation_rules.CLASSIFICATION_ROUTING:
            with self.subTest(classification=classification):
                rows = case_logic.query_playbook([classification], fake_query)
                self.assertTrue(rows)
        both = case_logic.query_playbook(["Tier downgrade", "Parity deviation"], fake_query)
        self.assertEqual(len(both), 3)  # Contracting Team + FRM for the downgrade, Contracting Team for parity


@unittest.skipUnless(yaml, "pip install pyyaml to check the index configuration")
class IndexConfiguration(unittest.TestCase):
    """The vector index can only filter on metadata keys NOT declared non-filterable, and filterable
    metadata is limited to about 2 KB per vector; long text must be declared non-filterable."""

    def test_long_playbook_fields_are_declared_non_filterable_and_the_rest_stays_small(self):
        class Loader(yaml.SafeLoader):
            pass
        Loader.add_multi_constructor("!", lambda l, s, n: l.construct_scalar(n) if isinstance(n, yaml.ScalarNode)
                                     else l.construct_sequence(n) if isinstance(n, yaml.SequenceNode) else l.construct_mapping(n))
        with open(os.path.join(ROOT, "template.yaml"), encoding="utf-8") as f:
            index = yaml.load(f, Loader=Loader)["Resources"]["VectorIndex"]["Properties"]
        non_filterable = set(index["MetadataConfiguration"]["NonFilterableMetadataKeys"])
        self.assertLessEqual(len(non_filterable), 10, "the service allows at most 10")
        docs, _ = ingest(XLSX)
        for doc in docs:
            filterable = {k: v for k, v in doc["metadata"].items() if k not in non_filterable}
            with self.subTest(row=doc["key"]):
                self.assertLess(len(json.dumps(filterable).encode()), 2048)
                for needed in ("docType", "classification", "persona", "approvalRequired"):
                    self.assertIn(needed, filterable, f"{needed} must stay filterable: retrieval filters on it")
        for long_field in ("nextBestAction", "requiredEvidence", "doNotDo", "triggerCondition", "sourceKey"):
            self.assertIn(long_field, non_filterable)


if __name__ == "__main__":
    unittest.main()
